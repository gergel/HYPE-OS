"""A régi (Excelben készült) árajánlatok kinyerése az árajánlat-készítőbe.

A felhasználó kérése (2026-10-29): a Google Drive „árajánlatok” mappájában lévő
összes régi árajánlat kerüljön fel az Árajánlatok oldalra „régi árajánlatként”,
a HYPE OS ajánlat-kinézetével, hogy elő lehessen venni, átírni és menteni.

Ez a szkript a letöltött XLSX-ekből állítja elő az
`app/quotes/regi_ajanlatok.json` adatfájlt, amit a `r1g2i3a4j5n6` migráció tölt
be a `quotes` / `quote_items` táblákba. Az XLSX-ek a megszokott HYPE-formát
követik: fejléc-sor (Tétel/Szolgáltatás | Alkalom | Mennyiség | Egységár |
Teljes ár), szekció-fejlécek (pl. „Emberi erőforrás”), összesítő sor
(„A PROJECT TELJES KÖLTSÉGE”), esetleg kedvezmény, végül „Megjegyzés:”.

Használat:

    python scripts/regi_arajanlatok_kinyerese.py LISTA.tsv LETOLTES_MAPPA [KIMENET.json]

LISTA.tsv soronként: drive_id <TAB> mappa (év) <TAB> módosítás ideje (ISO) <TAB> fájlnév;
a fájl a LETOLTES_MAPPA/<mappa>/<drive_id>.xlsx útvonalon van. A kimenet végén
a kihagyott fájlok listája (nem árajánlat-táblázat: ártábla, ajánlatkérő
melléklet stb.) a standard hibakimenetre kerül."""

from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ALAP_KIMENET = Path(__file__).resolve().parents[1] / "app" / "quotes" / "regi_ajanlatok.json"

NEV_FEJ = ("tétel", "item", "forgatás", "project", "emberi", "leírás")
ALKALOM_FEJ = ("alkalom", "nap", "szorzó", "unit", "project", "occasion")
MENNYISEG_FEJ = ("mennyiség", "quantity", "fő", "db", "óra", "szorzó")
EGYSEGAR_FEJ = ("egységár", "unit price", "nettó + áfa")
TELJES_FEJ = ("teljes ár", "total")
OSSZESITO = ("KÖLTSÉG", "VÉGÖSSZEG", "ÖSSZESEN", "TOTAL", "HAVIDÍJ", "PROJCT")
ELOTAGOK = ("HYPE_ÁRAJÁNLAT_", "HYPE_árajánlat_", "HYPE_Árajánlat_", "CB_ÁRAJÁNLAT_", "CB_Quote_",
            "LEDFAL-BERLES_ÁRAJÁNLAT_", "HYPE_QOUTATION_", "HYPE_ELSZÁMOLÁS_", "HYPE_ÁRLISTA_",
            "HYPE_Árajánlat ", "HYPE_")


def _nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s)


def _szoveg(v) -> str:
    return _nfc(str(v)).strip() if v is not None else ""


def _szam(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.replace("\xa0", "").replace(" ", "").replace("Ft", "").replace(",", ".")
        if re.fullmatch(r"-?\d+(\.\d+)?", s):
            return float(s)
    return None


def _fejlec_oszlopok(sor: list) -> dict | None:
    """Ha a sor fejléc (Tétel/Szolgáltatás | Alkalom | …), az oszlopindexek."""
    cimkek = [_szoveg(v).lower() for v in sor]
    o: dict = {}
    for i, c in enumerate(cimkek):
        if not c:
            continue
        if "nev" not in o and i <= 1 and (any(c.startswith(x) for x in NEV_FEJ) or i == 0):
            o.setdefault("nev_jelolt", i)
        if "alkalom" not in o and any(c.startswith(x) for x in ALKALOM_FEJ) and i > 0:
            o["alkalom"] = i
        elif "mennyiseg" not in o and any(c.startswith(x) for x in MENNYISEG_FEJ):
            o["mennyiseg"] = i
        elif "egysegar" not in o and any(c.startswith(x) for x in EGYSEGAR_FEJ):
            o["egysegar"] = i
        elif "teljes" not in o and any(c.startswith(x) for x in TELJES_FEJ):
            o["teljes"] = i
    if not ({"egysegar", "teljes"} & o.keys()) or not ({"alkalom", "mennyiseg", "teljes"} & o.keys()):
        return None
    if not ({"alkalom", "mennyiseg"} & o.keys()) and not any(x in cimkek[o.get("teljes", 0)] for x in TELJES_FEJ):
        return None
    adat_oszlopok = [o[k] for k in ("alkalom", "mennyiseg", "egysegar", "teljes") if k in o]
    o["nev"] = min(adat_oszlopok) - 1 if min(adat_oszlopok) > 0 else 0
    # A név oszlopa a legelső, nem üres cella a számoszlopok előtt.
    for i in range(0, min(adat_oszlopok)):
        if cimkek[i]:
            o["nev"] = i
            break
    o["fej_nev"] = _szoveg(sor[o["nev"]])
    o.pop("nev_jelolt", None)
    return o


def _nev_es_leiras(szoveg: str) -> tuple[str, str | None]:
    sorok = [s.rstrip() for s in szoveg.replace("\r", "").split("\n")]
    nev = sorok[0].strip()
    leiras = "\n".join(s for s in sorok[1:] if s.strip()).strip() or None
    if len(nev) > 255:
        leiras = (nev[255:] + ("\n" + leiras if leiras else "")).strip()
        nev = nev[:255]
    return nev, leiras


def _datum(cim: str) -> tuple[str | None, str | None]:
    m = re.search(r"(20\d\d)\.\s?(\d{1,2})\.\s?(\d{1,2})\.?(?:\s*[-–]\s*(?:(\d{1,2})\.\s?)?(\d{1,2}))?", cim)
    if not m:
        return None, None
    ev, ho, nap = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if not (1 <= ho <= 12 and 1 <= nap <= 31):
        return None, None
    tol = f"{ev:04d}-{ho:02d}-{nap:02d}"
    ig = None
    if m.group(5):
        ho2 = int(m.group(4)) if m.group(4) else ho
        nap2 = int(m.group(5))
        if 1 <= ho2 <= 12 and 1 <= nap2 <= 31 and (ho2, nap2) > (ho, nap):
            ig = f"{ev:04d}-{ho2:02d}-{nap2:02d}"
    return tol, ig


def munkalap(ws, ws_kepletek) -> dict | None:
    """Egy munkalap → ajánlat (szekciók, sorok, összesítő, kedvezmény,
    megjegyzés), vagy None, ha nem árajánlat-táblázat."""
    sorok = [list(r) for r in ws.iter_rows(values_only=True)]
    kepletek = [list(r) for r in ws_kepletek.iter_rows(values_only=True)]
    oszl = None
    cim_sorok: list[str] = []
    blokk: str | None = None
    szekcio: str | None = None
    tetelek: list[dict] = []
    osszesito_cimke = None
    osszesito_ertekek: list[float] = []
    kedv_szazalek = 0.0
    kedv_osszeg = 0
    megjegyzes = None
    csak_teljes = False
    kedv_sorok: list[str] = []
    utolso_osszesito: int | None = None
    osszesito_cimkek: list[str] = []

    def kovetkezo_tartalmas(i: int):
        for j in range(i + 1, len(sorok)):
            if any(v is not None and _szoveg(v) for v in sorok[j]):
                return sorok[j]
        return None

    def szamos(sor) -> bool:
        if oszl is None:
            return False
        return any(_szam(sor[oszl[k]]) is not None for k in ("alkalom", "mennyiseg", "egysegar", "teljes")
                   if k in oszl and oszl[k] < len(sor))

    for i, sor in enumerate(sorok):
        if not any(v is not None and _szoveg(v) for v in sor):
            continue
        fej = _fejlec_oszlopok(sor)
        if fej is not None:
            # A későbbi (szekció-)fejlécek felirata eltérhet („Fő”, „Nap”) - az
            # oszlopok helye az első fejlécé marad, ami hiányzik, azt pótoljuk.
            oszl = {**fej, **{k: v for k, v in (oszl or {}).items() if k != "fej_nev"}}
            csak_teljes = not ({"alkalom", "mennyiseg", "egysegar"} & oszl.keys())
            nev = fej["fej_nev"]
            if nev and not nev.lower().startswith(("tétel", "item")):
                szekcio = f"{blokk} – {nev}" if blokk else nev
            else:
                szekcio = blokk
            continue
        if oszl is None and len(sor) >= 4 and _szoveg(sor[0]) and all(_szam(sor[k]) is not None for k in (1, 2, 3)):
            # Fejléc nélküli, de a megszokott elrendezésű tábla (A | B | C | D | – | F).
            oszl = {"nev": 0, "alkalom": 1, "mennyiseg": 2, "egysegar": 3, "fej_nev": ""}
            if len(sor) > 5:
                oszl["teljes"] = 5
            if cim_sorok and blokk is None and len(cim_sorok) > 1:
                blokk = szekcio = cim_sorok.pop()[:120]
        if oszl is None:
            # A fejléc előtti sorok: cím (projekt neve, dátum) vagy blokk-cím.
            szoveg = " ".join(_szoveg(v) for v in sor if v is not None and _szoveg(v))
            if szoveg and _szam(szoveg) is None:
                cim_sorok.append(szoveg)
            continue
        nev_szoveg = _szoveg(sor[oszl["nev"]]) if oszl["nev"] < len(sor) else ""
        if not nev_szoveg:
            # Pl. a CHAINE-féle táblákban az első oszlop üres - nincs mit felvenni.
            continue
        nagy = nev_szoveg.upper()
        if nagy.startswith("MEGJEGYZÉS") or nagy.startswith("NOTE"):
            megjegyzes = re.sub(r"^(megjegyzés|note)s?\s*:?\s*", "", nev_szoveg, flags=re.I).strip() or None
            continue
        teljes = _szam(sor[oszl["teljes"]]) if "teljes" in oszl and oszl["teljes"] < len(sor) else None
        teljes_keplet = (
            "teljes" in oszl and i < len(kepletek) and oszl["teljes"] < len(kepletek[i])
            and isinstance(kepletek[i][oszl["teljes"]], str) and kepletek[i][oszl["teljes"]].startswith("=")
        )
        if "KEDVEZMÉNY" in nagy or "DISCOUNT" in nagy:
            kedv_sorok.append(nev_szoveg.split("*")[0].strip() + (f": {int(round(teljes))} Ft" if teljes else ""))
            m = re.search(r"(\d+(?:[.,]\d+)?)\s*%", nev_szoveg)
            if m:
                kedv_szazalek = float(m.group(1).replace(",", "."))
            elif teljes is not None:
                kedv_osszeg += abs(int(round(teljes)))
            continue
        alkalom = _szam(sor[oszl["alkalom"]]) if "alkalom" in oszl and oszl["alkalom"] < len(sor) else None
        mennyiseg = _szam(sor[oszl["mennyiseg"]]) if "mennyiseg" in oszl and oszl["mennyiseg"] < len(sor) else None
        egysegar = _szam(sor[oszl["egysegar"]]) if "egysegar" in oszl and oszl["egysegar"] < len(sor) else None
        if alkalom is None and mennyiseg is None and egysegar is None and (teljes is not None or teljes_keplet) and any(
            x in nagy for x in OSSZESITO
        ):
            cimke = re.split(r"\*", nev_szoveg)[0].strip().rstrip(":").strip()
            if osszesito_cimke is None and "VÉGÖSSZEG" not in nagy and "HAVIDÍJ" not in nagy:
                osszesito_cimke = cimke[:120]
            osszesito_cimkek.append(cimke)
            utolso_osszesito = len(tetelek)
            if "HAVIDÍJ" not in nagy and teljes is not None:
                osszesito_ertekek.append(teljes)
            continue
        if not szamos(sor):
            # Számok nélküli sor: szekció-cím, leírás-folytatás vagy 0 Ft-os tétel.
            if "\n" in nev_szoveg.strip():
                nev, leiras = _nev_es_leiras(nev_szoveg)
                tetelek.append({"section": szekcio, "name": nev, "description": leiras, "occasions": 1,
                                "quantity": 1, "unit_price": 0})
                continue
            kov = kovetkezo_tartalmas(i)
            kov_szamos = kov is not None and (szamos(kov) or _fejlec_oszlopok(kov) is not None)
            if kov_szamos or not tetelek:
                if kov is not None and _fejlec_oszlopok(kov) is not None:
                    blokk = nev_szoveg[:120]
                    szekcio = blokk
                else:
                    szekcio = f"{blokk} – {nev_szoveg}"[:120] if blokk and blokk != nev_szoveg else nev_szoveg[:120]
            else:
                elozo = tetelek[-1]
                elozo["description"] = ((elozo["description"] + "\n") if elozo["description"] else "") + nev_szoveg
            continue
        if csak_teljes:
            alkalom, mennyiseg, egysegar = 1.0, 1.0, teljes or 0.0
        else:
            alkalom = 1.0 if alkalom is None else alkalom
            mennyiseg = 1.0 if mennyiseg is None else mennyiseg
            if egysegar is None:
                # Csak teljes ár: az egységár abból (alkalom × mennyiség osztva).
                szorzo = alkalom * mennyiseg
                egysegar = (teljes / szorzo) if (teljes is not None and szorzo) else (teljes or 0.0)
        nev, leiras = _nev_es_leiras(nev_szoveg)
        sorkedvezmeny = 0
        teljes_szoveg = _szoveg(sor[oszl["teljes"]]) if "teljes" in oszl and oszl["teljes"] < len(sor) else ""
        if teljes is None and teljes_szoveg and not teljes_keplet and _szam(teljes_szoveg) is None:
            # Pl. „Keretszerződés részeként kerül elszámolásra”: az ár látszik,
            # de a végösszegbe nem számít bele (100% sorkedvezmény).
            sorkedvezmeny = 100
            leiras = ((leiras + "\n") if leiras else "") + f"({teljes_szoveg})"
        elif teljes is not None and not csak_teljes and egysegar > 0 and alkalom * mennyiseg > 0:
            # Kézzel átírt teljes ár (pl. 20 000 helyett 15 000): sorkedvezményként.
            vart = egysegar * alkalom * mennyiseg
            if 0 <= teljes < vart - 1:
                sorkedvezmeny = round((1 - teljes / vart) * 100, 2)
        if egysegar < 0:
            kedv_osszeg += int(round(-egysegar * alkalom * mennyiseg))
            continue
        tetelek.append({
            "section": szekcio[:120] if szekcio else None, "name": nev, "description": leiras,
            "occasions": round(alkalom, 2), "quantity": round(mennyiseg, 2), "unit_price": int(round(egysegar)),
            **({"line_discount_percent": sorkedvezmeny} if sorkedvezmeny else {}),
        })

    if oszl is None or not tetelek or not any(t["unit_price"] for t in tetelek):
        return None
    if utolso_osszesito is not None:
        # Az utolsó összesítő UTÁNI tételek (pl. „Utómunka árlista”, „Utólagos
        # elszámolást igényel”) nem részei a végösszegnek: opcionális sorok.
        for t in tetelek[utolso_osszesito:]:
            t["is_optional"] = True
    cim = " · ".join(dict.fromkeys(c for c in cim_sorok if c))
    megjegyzes_belso = None
    nem_vegosszeg = [c for c in osszesito_cimkek if "VÉGÖSSZEG" not in c.upper()]
    if len(nem_vegosszeg) > 1 and kedv_sorok:
        # Több forgatókönyves árlista (pl. „1 nap” / „20 nap kedvezménnyel”):
        # a kedvezmény csak egy forgatókönyvre vonatkozott - nem vonjuk le.
        kedv_szazalek, kedv_osszeg = 0.0, 0
        megjegyzes_belso = "Az eredeti táblázat kedvezményei: " + "; ".join(kedv_sorok)
    return {
        "cim": cim, "items": tetelek, "summary_label": osszesito_cimke,
        "discount_percent": min(kedv_szazalek, 100.0), "discount_amount": kedv_osszeg, "note_text": megjegyzes,
        "eredeti_osszegek": osszesito_ertekek, "megjegyzes_belso": megjegyzes_belso,
    }


def _projekt_es_ugyfel(fajlnev: str) -> tuple[str, str, str | None]:
    """(márka, projekt neve, ügyfél-jelölt) a fájlnévből."""
    torzs = _nfc(fajlnev)
    torzs = re.sub(r"\.xlsx$", "", torzs, flags=re.I).strip()
    brand = "CB" if torzs.upper().startswith("CB_") else "HYPE"
    for e in ELOTAGOK:
        if torzs.upper().startswith(_nfc(e).upper()):
            torzs = torzs[len(e):]
            break
    torzs = torzs.strip(" _-")
    ugyfel = None
    m = re.match(r"^(.+?)\s+-\s+", torzs) or re.match(r"^([A-ZÁÉÍÓÖŐÚÜŰ]{2,})_", torzs)
    if m:
        ugyfel = m.group(1).strip()
    return brand, torzs[:255], ugyfel


def kinyer(lista: Path, mappa: Path) -> tuple[list[dict], list[str]]:
    ki: list[dict] = []
    kihagyott: list[str] = []
    latott: set[str] = set()
    for sor in lista.read_text(encoding="utf-8").splitlines():
        drive_id, ev, mikor, fajlnev = sor.split("\t")
        fajlnev = _nfc(fajlnev)
        if fajlnev.startswith("~$"):
            continue
        ut = mappa / ev / f"{drive_id}.xlsx"
        try:
            wb = openpyxl.load_workbook(ut, data_only=True)
            wbk = openpyxl.load_workbook(ut, data_only=False)
        except Exception as exc:  # noqa: BLE001 - sérült / nem xlsx fájl
            kihagyott.append(f"{ev}/{fajlnev}: nem olvasható ({exc})")
            continue
        lapok = []
        for ws in wb.worksheets:
            try:
                a = munkalap(ws, wbk[ws.title])
            except Exception as exc:  # noqa: BLE001
                kihagyott.append(f"{ev}/{fajlnev} [{ws.title}]: hiba ({exc})")
                continue
            if a is not None:
                lapok.append((ws.title, a))
        if not lapok:
            kihagyott.append(f"{ev}/{fajlnev}: nem árajánlat-táblázat")
            continue
        brand, projekt, ugyfel = _projekt_es_ugyfel(fajlnev)
        for lap_nev, a in lapok:
            kulcs_alap = json.dumps(a["items"], ensure_ascii=False, sort_keys=True) + projekt
            ujjlenyomat = hashlib.sha1(kulcs_alap.encode()).hexdigest()
            if ujjlenyomat in latott:
                continue
            latott.add(ujjlenyomat)
            nev = projekt if len(lapok) == 1 else f"{projekt} ({lap_nev})"[:255]
            tol, ig = _datum(a["cim"])
            penznem_eur = "EUR" in fajlnev.upper() or "QUOTE" in fajlnev.upper() or "QOUTATION" in fajlnev.upper()
            ki.append({
                "source": f"{ev}/{fajlnev}" + (f" [{lap_nev}]" if len(lapok) > 1 else ""),
                "drive_id": drive_id,
                "date": mikor,
                "year": int(ev) if ev.isdigit() else int(mikor[:4]),
                "brand": brand,
                "project_name": nev,
                "client_guess": ugyfel,
                "title": a["cim"] or None,
                "event_date_from": tol,
                "event_date_to": ig,
                "summary_label": a["summary_label"],
                "discount_percent": a["discount_percent"],
                "discount_amount": a["discount_amount"],
                "note_text": a["note_text"],
                "currency_note": "Az eredeti ajánlat euróban (EUR) készült." if penznem_eur else None,
                "original_totals": a["eredeti_osszegek"],
                "internal_note": a["megjegyzes_belso"],
                "items": a["items"],
            })
    ki.sort(key=lambda x: (x["date"], x["source"]))
    return ki, kihagyott


def main() -> None:
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    kimenet = Path(sys.argv[3]) if len(sys.argv) > 3 else ALAP_KIMENET
    ajanlatok, kihagyott = kinyer(Path(sys.argv[1]), Path(sys.argv[2]))
    kimenet.write_text(json.dumps(ajanlatok, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{len(ajanlatok)} ajánlat → {kimenet}")
    for k in kihagyott:
        print("KIHAGYVA:", k, file=sys.stderr)


if __name__ == "__main__":
    main()
