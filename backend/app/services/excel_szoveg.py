"""Excel/CSV melléklet SZÖVEGGÉ alakítása a részletező-kiolvasáshoz (lásd
services/szamla_erkeztetes._reszletezo_feldolgozas): a bejövő számlák mellé
érkező költség-részletező táblázatokat (projektbontás) nem képként, hanem
cellaértékekként adjuk át a Gemininek - így a számok pontosan, kerekítés és
OCR-hiba nélkül jutnak el a kiolvasásig."""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

#: A részletezők pár tucat sorosak - egy ezres sor-plafon bőven elég, és
#: megvéd attól, hogy egy óriási exportált táblázat token-tengerré váljon.
MAX_SOR = 1000
MAX_OSZLOP = 40

EXCEL_MIME = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",  # .xlsx
    "application/vnd.ms-excel",  # .xls (a Gmail néha az .xlsx-re is ezt adja)
    "text/csv",
    "application/csv",
}

#: Fájlnév-végződések, amikre Excel-ként próbálkozunk akkor is, ha a MIME
#: általánosabb (pl. application/octet-stream) - a levelezők gyakran ezt adják.
EXCEL_KITERJESZTESEK = (".xlsx", ".xlsm", ".csv")


def excelnek_tunik(mime_type: str | None, fajlnev: str | None) -> bool:
    """Táblázat-melléklet-e - MIME vagy kiterjesztés alapján."""
    if (mime_type or "").lower() in EXCEL_MIME:
        return True
    return (fajlnev or "").lower().endswith(EXCEL_KITERJESZTESEK)


def szovegge(adat: bytes, fajlnev: str | None = None) -> str:
    """A táblázat tartalma tabulátorral tagolt szövegként, munkalaponként.

    ValueError-t dob emberi hibaüzenettel, ha a fájl nem nyitható meg
    táblázatként (sérült, jelszavas, vagy valójában nem is Excel)."""
    nev = (fajlnev or "").lower()
    if nev.endswith(".csv") or adat[:4] not in (b"PK\x03\x04",):
        # Nem zip-alapú (xlsx) fájl: CSV-ként próbáljuk. A csv-ág elé vett
        # zip-ellenőrzés azért kell, mert a .xls/.csv MIME-ok keverednek.
        if not nev.endswith((".xlsx", ".xlsm")):
            szoveg = _csv_szovegge(adat)
            if szoveg is not None:
                return szoveg
    try:
        munkafuzet = load_workbook(io.BytesIO(adat), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - a hívó emberi hibaüzenetet vár
        raise ValueError(f"A táblázat nem nyitható meg ({exc}).") from exc

    reszek: list[str] = []
    try:
        for lap in munkafuzet.worksheets:
            sorok: list[str] = []
            for i, sor in enumerate(lap.iter_rows(values_only=True)):
                if i >= MAX_SOR:
                    sorok.append(f"... (a munkalap további sorai levágva, {MAX_SOR} sor felett)")
                    break
                ertekek = ["" if c is None else str(c) for c in sor[:MAX_OSZLOP]]
                if any(e.strip() for e in ertekek):
                    sorok.append("\t".join(ertekek).rstrip())
            if sorok:
                reszek.append(f"=== Munkalap: {lap.title} ===\n" + "\n".join(sorok))
    finally:
        munkafuzet.close()

    if not reszek:
        raise ValueError("A táblázat üres - nincs kiolvasható cella.")
    return "\n\n".join(reszek)


def _csv_szovegge(adat: bytes) -> str | None:
    """CSV-tartalom tabulátoros szöveggé; None, ha nem értelmezhető CSV-ként."""
    for kodolas in ("utf-8-sig", "utf-8", "cp1250", "latin-1"):
        try:
            szoveg = adat.decode(kodolas)
            break
        except UnicodeDecodeError:
            continue
    else:
        return None
    try:
        elvalaszto = csv.Sniffer().sniff(szoveg[:4096], delimiters=",;\t").delimiter
    except csv.Error:
        elvalaszto = ";" if szoveg.count(";") >= szoveg.count(",") else ","
    sorok: list[str] = []
    for i, sor in enumerate(csv.reader(io.StringIO(szoveg), delimiter=elvalaszto)):
        if i >= MAX_SOR:
            sorok.append(f"... (további sorok levágva, {MAX_SOR} sor felett)")
            break
        ertekek = [c.strip() for c in sor[:MAX_OSZLOP]]
        if any(ertekek):
            sorok.append("\t".join(ertekek).rstrip())
    return "\n".join(sorok) if sorok else None
