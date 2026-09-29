"""Árajánlat export: XLSX (a megszokott HYPE Excel-formában, képletekkel) és
PDF (ugyanabból az adatból).

Felépítés (mindkettőben):
    [HYPE vagy ContentBee logó]                       cégadatok, ajánlatszám
    PROJECT NEVE: <projekt neve>
    Tétel/Szolgáltatás | Alkalom | Mennyiség | Egységár | Teljes ár
    <szekció fejléc>
    <sorok…>  (a leírás a tétel cellájában új sorokban, „- ” felsorolással)
    <CÍMKE>   *Árak az ÁFÁt NEM tartalmazzák.            <nettó végösszeg>
    (havidíjas: A PROJECT HAVIDÍJA / A PROJECT TELJES KÖLTSÉGE N HÓNAPRA)
    Opcionális tételek (külön blokk, saját összeggel)
    Megjegyzés: <megjegyzés>

Az XLSX „Teljes ár” oszlopa képlet (=B*C*D), a végösszeg =SUM(...)."""

from __future__ import annotations

import io
import re
from datetime import date
from pathlib import Path

from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.client import Client
from app.models.quote import Quote, QuoteItem
from app.quotes.service import ALAP_CIMKE, osszegzes
from app.quotes.szamitas import sor_osszeg

ASSETS = Path(__file__).parent / "assets"
LOGO = {"HYPE": ASSETS / "logo-hype.png", "CB": ASSETS / "logo-cb.png"}
FAJL_ELOTAG = {"HYPE": "HYPE", "CB": "CB"}
AFA_MEGJEGYZES = "*Árak az ÁFÁt NEM tartalmazzák."
OSZLOPOK = ["Tétel/Szolgáltatás", "Alkalom", "Mennyiség", "Egységár", "Teljes ár"]


def _ceg() -> list[str]:
    return [settings.arajanlat_ceg_nev, settings.arajanlat_ceg_cim, f"Adószám: {settings.arajanlat_ceg_adoszam}"]


def _ugyfel(db: Session, q: Quote) -> str | None:
    c = db.get(Client, q.client_id) if q.client_id else None
    return c.nev if c else None


def fajlnev(db: Session, q: Quote, kiterjesztes: str) -> str:
    """HYPE_ÁRAJÁNLAT_<ÜGYFÉL> - <PROJEKT NEVE>.xlsx (nagybetűvel, mint eddig)."""
    reszek = [x for x in (_ugyfel(db, q), q.project_name) if x and x.strip()]
    torzs = " - ".join(reszek) or q.number
    if q.version > 1:
        torzs += f" V{q.version}"
    torzs = re.sub(r'[\\/:*?"<>|\r\n]+', " ", torzs).strip().upper()
    return f"{FAJL_ELOTAG.get(q.brand, 'HYPE')}_ÁRAJÁNLAT_{torzs}.{kiterjesztes}"


def _szam(x) -> float | int:
    f = float(x or 0)
    return int(f) if f.is_integer() else f


def _datumok(q: Quote) -> str | None:
    if not q.event_date_from:
        return None
    fmt = lambda d: d.strftime("%Y.%m.%d.")  # noqa: E731
    if q.event_date_to and q.event_date_to != q.event_date_from:
        return f"{fmt(q.event_date_from)} – {fmt(q.event_date_to)}"
    return fmt(q.event_date_from)


def _csoportok(sorok: list[QuoteItem]) -> list[tuple[str | None, list[QuoteItem]]]:
    """Egymás után következő, azonos szekciójú sorok egy csoportban."""
    ki: list[tuple[str | None, list[QuoteItem]]] = []
    for it in sorted(sorok, key=lambda x: x.sort_order):
        sz = (it.section or "").strip() or None
        if ki and ki[-1][0] == sz:
            ki[-1][1].append(it)
        else:
            ki.append((sz, [it]))
    return ki


def _osszesito_cimkek(q: Quote) -> tuple[str, str | None]:
    """(a végösszeg sor címkéje, a havidíj sor címkéje vagy None)."""
    if q.pricing_mode == "monthly":
        return f"A PROJECT TELJES KÖLTSÉGE {q.months} HÓNAPRA", "A PROJECT HAVIDÍJA"
    return (q.summary_label or ALAP_CIMKE).strip().upper(), None


def _teljes_ar_fejlec(q: Quote) -> str:
    return f"Teljes ár ({q.months} hónap)" if q.pricing_mode == "monthly" else "Teljes ár"


# ── XLSX ─────────────────────────────────────────────────────────────────────


def xlsx(db: Session, q: Quote) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Árajánlat"
    for oszlop, szel in zip("ABCDE", (62, 10, 11, 15, 17)):
        ws.column_dimensions[oszlop].width = szel
    PENZ = '#,##0 "Ft"'
    SZAM = "0.##"
    vekony = Side(style="thin", color="D0D0D0")
    keret = Border(bottom=vekony)
    sotet = PatternFill("solid", fgColor="1F1F1F")
    halvany = PatternFill("solid", fgColor="EFEFEF")
    fonts = {"cim": Font(bold=True, size=14), "fej": Font(bold=True, color="FFFFFF"),
             "szekcio": Font(bold=True), "osszeg": Font(bold=True, size=12), "kicsi": Font(size=9, color="555555"),
             "dolt": Font(italic=True, size=9, color="555555")}

    # Logó + cégadatok, ajánlatszám.
    logo = LOGO.get(q.brand)
    if logo and logo.exists():
        img = XlImage(str(logo))
        arany = 70 / img.height
        img.height, img.width = 70, int(img.width * arany)
        ws.add_image(img, "A1")
    for i, sor in enumerate(_ceg() + [f"Ajánlatszám: {q.number}" + (f" (v{q.version})" if q.version > 1 else ""),
                                      f"Kelt: {date.today().strftime('%Y.%m.%d.')}"]):
        c = ws.cell(row=1 + i, column=4, value=sor)
        c.font = fonts["kicsi"]
    ws.row_dimensions[1].height = 15

    r = 7
    c = ws.cell(row=r, column=1, value=f"PROJECT NEVE: {q.project_name}")
    c.font = fonts["cim"]
    info = [x for x in (
        f"Megrendelő: {_ugyfel(db, q)}" if _ugyfel(db, q) else None,
        f"Időpont: {_datumok(q)}" if _datumok(q) else None,
        f"Helyszín: {q.location}" if q.location else None,
    ) if x]
    for x in info:
        r += 1
        ws.cell(row=r, column=1, value=x).font = fonts["kicsi"]
    r += 2

    fejlec = OSZLOPOK[:4] + [_teljes_ar_fejlec(q)]
    for j, nev in enumerate(fejlec, start=1):
        c = ws.cell(row=r, column=j, value=nev)
        c.font, c.fill = fonts["fej"], sotet
        c.alignment = Alignment(horizontal="left" if j == 1 else "center", vertical="center")
    r += 1

    def sorok_ki(sorok: list[QuoteItem], kezd: int) -> tuple[int, list[int]]:
        sor = kezd
        tetel_sorok: list[int] = []
        for szekcio, csoport in _csoportok(sorok):
            if szekcio:
                c = ws.cell(row=sor, column=1, value=szekcio)
                c.font = fonts["szekcio"]
                for j in range(1, 6):
                    ws.cell(row=sor, column=j).fill = halvany
                sor += 1
            for it in csoport:
                leiras = (it.description or "").strip()
                kedv = float(it.line_discount_percent or 0)
                if kedv:
                    leiras = (leiras + "\n" if leiras else "") + f"({_szam(kedv)}% kedvezménnyel)"
                if leiras:
                    ertek = CellRichText(TextBlock(InlineFont(b=True), it.name or ""), "\n" + leiras)
                else:
                    ertek = it.name or ""
                a = ws.cell(row=sor, column=1, value=ertek)
                if not leiras:
                    a.font = Font(bold=True)
                a.alignment = Alignment(wrap_text=True, vertical="top")
                ws.cell(row=sor, column=2, value=_szam(it.occasions)).number_format = SZAM
                ws.cell(row=sor, column=3, value=_szam(it.quantity)).number_format = SZAM
                ws.cell(row=sor, column=4, value=int(it.unit_price or 0)).number_format = PENZ
                keplet = f"=B{sor}*C{sor}*D{sor}" if not kedv else f"=ROUND(B{sor}*C{sor}*D{sor}*(1-{_szam(kedv)}/100),0)"
                ws.cell(row=sor, column=5, value=keplet).number_format = PENZ
                for j in range(2, 6):
                    ws.cell(row=sor, column=j).alignment = Alignment(horizontal="center" if j < 4 else "right",
                                                                     vertical="top")
                for j in range(1, 6):
                    ws.cell(row=sor, column=j).border = keret
                sorok_szama = 1 + leiras.count("\n") + (1 if leiras else 0) + len(leiras) // 90
                ws.row_dimensions[sor].height = max(18, 15 * sorok_szama)
                tetel_sorok.append(sor)
                sor += 1
        return sor, tetel_sorok

    alap_sorok = [i for i in q.items if not i.is_optional]
    opcios = [i for i in q.items if i.is_optional]
    r, tetel_sorok = sorok_ki(alap_sorok, r)
    osszeg_keplet = f"=SUM({','.join(f'E{x}' for x in tetel_sorok)})" if tetel_sorok else "=0"
    o = osszegzes(q)

    r += 1
    vegosszeg_cella = f"E{r}"
    if o.kedvezmeny:
        ws.cell(row=r, column=1, value="Részösszeg").font = fonts["szekcio"]
        ws.cell(row=r, column=5, value=osszeg_keplet).number_format = PENZ
        resz = f"E{r}"
        r += 1
        pct = float(q.discount_percent or 0)
        cimke = "Kedvezmény" + (f" ({_szam(pct)}%)" if pct else "")
        keplet = "=-(" + (f"ROUND({resz}*{_szam(pct)}/100,0)" if pct else "0") + f"+{int(q.discount_amount or 0)})"
        ws.cell(row=r, column=1, value=cimke).font = fonts["szekcio"]
        ws.cell(row=r, column=5, value=f"=MAX({keplet[1:]},-{resz})").number_format = PENZ
        kedv_cella = f"E{r}"
        r += 1
        osszeg_keplet = f"={resz}+{kedv_cella}"
        vegosszeg_cella = f"E{r}"

    cimke, havi_cimke = _osszesito_cimkek(q)
    c = ws.cell(row=r, column=1, value=cimke)
    c.font = fonts["osszeg"]
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    ws.cell(row=r, column=2, value=AFA_MEGJEGYZES).font = fonts["dolt"]
    c = ws.cell(row=r, column=5, value=osszeg_keplet)
    c.font, c.number_format = fonts["osszeg"], PENZ
    for j in range(1, 6):
        ws.cell(row=r, column=j).border = Border(top=Side(style="medium"), bottom=Side(style="medium"))
    if havi_cimke:
        r += 1
        ws.cell(row=r, column=1, value=havi_cimke).font = fonts["osszeg"]
        c = ws.cell(row=r, column=5, value=f"=ROUND({vegosszeg_cella}/{q.months},0)")
        c.font, c.number_format = fonts["osszeg"], PENZ

    if opcios:
        r += 2
        c = ws.cell(row=r, column=1, value="Opcionális tételek")
        c.font = fonts["szekcio"]
        for j in range(1, 6):
            ws.cell(row=r, column=j).fill = halvany
        r += 1
        r, opcio_sorok = sorok_ki(opcios, r)
        ws.cell(row=r, column=1, value="Opcionális tételek összesen").font = fonts["szekcio"]
        c = ws.cell(row=r, column=5, value=f"=SUM({','.join(f'E{x}' for x in opcio_sorok)})")
        c.font, c.number_format = fonts["szekcio"], PENZ
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.cell(row=r, column=2, value="(nem része a végösszegnek)").font = fonts["dolt"]

    if (q.note_text or "").strip():
        r += 2
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c = ws.cell(row=r, column=1, value="Megjegyzés:\n" + q.note_text.strip())
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(size=9)
        sorok = sum(1 + len(s) // 120 for s in ("Megjegyzés:\n" + q.note_text.strip()).split("\n"))
        ws.row_dimensions[r].height = 12.5 * sorok

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────

_BETUK_REGISZTRALVA = False


def _betuk() -> None:
    global _BETUK_REGISZTRALVA
    if _BETUK_REGISZTRALVA:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("AJ", str(ASSETS / "fonts" / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("AJ-B", str(ASSETS / "fonts" / "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFontFamily("AJ", normal="AJ", bold="AJ-B", italic="AJ", boldItalic="AJ-B")
    _BETUK_REGISZTRALVA = True


def _ft(n: int) -> str:
    return f"{int(n):,}".replace(",", " ") + " Ft"


def _esc(s: str) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def pdf(db: Session, q: Quote) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _betuk()
    alap = ParagraphStyle("alap", fontName="AJ", fontSize=8.8, leading=11.5)
    kicsi = ParagraphStyle("kicsi", parent=alap, fontSize=7.8, leading=10, textColor=colors.HexColor("#555555"))
    jobb = ParagraphStyle("jobb", parent=kicsi, alignment=TA_RIGHT)
    cim = ParagraphStyle("cim", parent=alap, fontName="AJ-B", fontSize=13, leading=16)
    fej = ParagraphStyle("fej", parent=alap, fontName="AJ-B", fontSize=8.2, textColor=colors.white)
    fej_j = ParagraphStyle("fejj", parent=fej, alignment=TA_RIGHT)
    vastag = ParagraphStyle("vastag", parent=alap, fontName="AJ-B")
    szam = ParagraphStyle("szam", parent=alap, alignment=TA_RIGHT)
    szam_b = ParagraphStyle("szamb", parent=szam, fontName="AJ-B", fontSize=10.5, leading=13)
    osszeg_cim = ParagraphStyle("osszegcim", parent=vastag, fontSize=10.5, leading=13)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm,
                            bottomMargin=14 * mm, title=fajlnev(db, q, "pdf")[:-4], author=settings.arajanlat_ceg_nev)
    szel = doc.width
    elemek: list = []

    logo = LOGO.get(q.brand)
    bal = Image(str(logo), height=20 * mm, width=20 * mm * _kep_arany(logo)) if logo and logo.exists() else Spacer(1, 1)
    ceg = "<br/>".join(_esc(x) for x in _ceg() + [
        f"Ajánlatszám: {q.number}" + (f" (v{q.version})" if q.version > 1 else ""),
        f"Kelt: {date.today().strftime('%Y.%m.%d.')}"])
    fejlec = Table([[bal, Paragraph(ceg, jobb)]], colWidths=[szel * 0.5, szel * 0.5])
    fejlec.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    elemek += [fejlec, Spacer(1, 6 * mm), Paragraph(_esc(f"PROJECT NEVE: {q.project_name}"), cim)]
    for x in (
        f"Megrendelő: {_ugyfel(db, q)}" if _ugyfel(db, q) else None,
        f"Időpont: {_datumok(q)}" if _datumok(q) else None,
        f"Helyszín: {q.location}" if q.location else None,
    ):
        if x:
            elemek.append(Paragraph(_esc(x), kicsi))
    elemek.append(Spacer(1, 4 * mm))

    oszl = [szel * 0.43, szel * 0.11, szel * 0.13, szel * 0.15, szel * 0.18]

    def tabla(sorok: list[QuoteItem], fejjel: bool) -> tuple[list, list]:
        adat: list = []
        stilus: list = []
        if fejjel:
            adat.append([Paragraph(OSZLOPOK[0], fej), Paragraph(OSZLOPOK[1], fej_j), Paragraph(OSZLOPOK[2], fej_j),
                         Paragraph(OSZLOPOK[3], fej_j), Paragraph(_teljes_ar_fejlec(q), fej_j)])
            stilus.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F1F1F")))
        for szekcio, csoport in _csoportok(sorok):
            if szekcio:
                i = len(adat)
                adat.append([Paragraph(_esc(szekcio), vastag), "", "", "", ""])
                stilus += [("SPAN", (0, i), (-1, i)), ("BACKGROUND", (0, i), (-1, i), colors.HexColor("#EFEFEF"))]
            for it in csoport:
                leiras = _esc((it.description or "").strip()).replace("\n", "<br/>")
                kedv = float(it.line_discount_percent or 0)
                if kedv:
                    leiras += ("<br/>" if leiras else "") + f"({_szam(kedv)}% kedvezménnyel)"
                szoveg = f"<b>{_esc(it.name)}</b>" + (f"<br/><font size=7.8 color='#444444'>{leiras}</font>" if leiras else "")
                adat.append([
                    Paragraph(szoveg, alap), Paragraph(str(_szam(it.occasions)).replace(".", ","), szam),
                    Paragraph(str(_szam(it.quantity)).replace(".", ","), szam), Paragraph(_ft(it.unit_price), szam),
                    Paragraph(_ft(sor_osszeg(it.occasions, it.quantity, it.unit_price, it.line_discount_percent)), szam),
                ])
                stilus.append(("LINEBELOW", (0, len(adat) - 1), (-1, len(adat) - 1), 0.4, colors.HexColor("#D0D0D0")))
        return adat, stilus

    o = osszegzes(q)
    adat, stilus = tabla([i for i in q.items if not i.is_optional], True)
    if o.kedvezmeny:
        pct = float(q.discount_percent or 0)
        adat.append([Paragraph("Részösszeg", vastag), "", "", "", Paragraph(_ft(o.reszosszeg), szam)])
        adat.append([Paragraph("Kedvezmény" + (f" ({_szam(pct)}%)" if pct else ""), vastag), "", "", "",
                     Paragraph("−" + _ft(o.kedvezmeny), szam)])
    cimke, havi_cimke = _osszesito_cimkek(q)
    i = len(adat)
    adat.append([Paragraph(f"{_esc(cimke)}<br/><font name='AJ' size=7.5 color='#555555'>{AFA_MEGJEGYZES}</font>",
                           osszeg_cim), "", "", "", Paragraph(_ft(o.netto), szam_b)])
    stilus += [("SPAN", (0, i), (3, i)), ("LINEABOVE", (0, i), (-1, i), 1.2, colors.black),
               ("LINEBELOW", (0, i), (-1, i), 1.2, colors.black)]
    if havi_cimke:
        adat.append([Paragraph(havi_cimke, osszeg_cim), "", "", "", Paragraph(_ft(o.havidij or 0), szam_b)])
        stilus.append(("SPAN", (0, len(adat) - 1), (3, len(adat) - 1)))
    t = Table(adat, colWidths=oszl, repeatRows=1)
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 3), *stilus]))
    elemek.append(t)

    opcios = [i for i in q.items if i.is_optional]
    if opcios:
        elemek += [Spacer(1, 5 * mm), Paragraph("Opcionális tételek", vastag), Spacer(1, 1.5 * mm)]
        adat, stilus = tabla(opcios, False)
        adat.append([Paragraph("Opcionális tételek összesen", vastag), Paragraph("(nem része a végösszegnek)", kicsi),
                     "", "", Paragraph(_ft(o.opcionalis), szam)])
        stilus += [("SPAN", (1, len(adat) - 1), (3, len(adat) - 1)),
                   ("LINEABOVE", (0, len(adat) - 1), (-1, len(adat) - 1), 0.8, colors.black)]
        t = Table(adat, colWidths=oszl)
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), *stilus]))
        elemek.append(t)

    if (q.note_text or "").strip():
        elemek += [Spacer(1, 6 * mm), Paragraph("<b>Megjegyzés:</b>", kicsi)]
        for bek in q.note_text.strip().split("\n"):
            elemek.append(Paragraph(_esc(bek), kicsi))

    doc.build(elemek)
    return buf.getvalue()


def _kep_arany(p: Path) -> float:
    from PIL import Image as PILImage

    with PILImage.open(p) as im:
        return im.width / im.height
