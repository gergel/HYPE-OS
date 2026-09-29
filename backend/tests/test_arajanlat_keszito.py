"""Árajánlat-készítő modul (2026-09): összegzés, számozás, seed, sablonból
létrehozás, ár-snapshot, variáns / verzió, sablon mentése, export, API.

Az összegzés tesztjei tiszták (DB nélkül); a többi Postgres-integráció EGY
tranzakcióban, a végén rollback (az API-hívásoknál a commit flush-ra cserélve)."""

from __future__ import annotations

import io
from decimal import ROUND_HALF_UP, Decimal
from types import SimpleNamespace

import pytest
from sqlalchemy import select
from sqlalchemy.exc import OperationalError

from app.quotes.szamitas import osszesit, sor_osszeg


def _sor(alkalom=1, menny=1, ar=0, kedv=0, opcio=False):
    return SimpleNamespace(occasions=alkalom, quantity=menny, unit_price=ar, line_discount_percent=kedv,
                           is_optional=opcio)


# ── Összegzés (tiszta) ───────────────────────────────────────────────────────


def test_sor_osszeg_szorzat_es_sorkedvezmeny():
    assert sor_osszeg(2, 3, 80000) == 480000
    assert sor_osszeg(1, 1, 85000, 10) == 76500
    assert sor_osszeg(1.5, 1, 10001) == 15002  # forintra kerekít (15001,5 → 15002)
    assert sor_osszeg(1, 0, 150) == 0


def test_opcionalis_sor_nem_szamit_bele_kulon_osszeg():
    o = osszesit([_sor(ar=100000), _sor(ar=50000, opcio=True), _sor(1, 2, 7500)])
    assert o.reszosszeg == 115000 and o.opcionalis == 50000 and o.netto == 115000
    assert o.afa == 31050 and o.brutto == 146050 and o.havidij is None


def test_ajanlat_kedvezmeny_szazalek_es_osszeg_legfeljebb_a_reszosszegig():
    o = osszesit([_sor(ar=200000)], kedvezmeny_szazalek=10, kedvezmeny_osszeg=5000)
    assert o.kedvezmeny == 25000 and o.netto == 175000
    assert osszesit([_sor(ar=1000)], kedvezmeny_osszeg=5000).netto == 0


def test_havidijas_mod_havidij_a_vegosszeg_per_honap():
    o = osszesit([_sor(1, 12, 300000), _sor(1, 12, 90000)], arazas="monthly", honapok=12)
    assert o.netto == 4680000 and o.havidij == 390000
    assert osszesit([_sor(ar=100)], arazas="monthly", honapok=3).havidij == 33


# ── Adatbázisos tesztek ──────────────────────────────────────────────────────


@pytest.fixture()
def db():
    from app.core.database import SessionLocal

    try:
        sess = SessionLocal()
        sess.execute(select(1))
    except OperationalError:
        pytest.skip("Postgres nem elérhető — integrációs teszt kihagyva.")
    try:
        yield sess
    finally:
        sess.rollback()
        sess.close()


@pytest.fixture()
def admin(db):
    from app.models.employee import Employee

    e = db.get(Employee, 2)
    if e is None:
        pytest.skip("Nincs admin munkatárs (#2) a teszthez.")
    return e


@pytest.fixture()
def alap(db):
    from app.quotes.seed import seed

    seed(db)
    return db


@pytest.fixture()
def ugyfel(db):
    from app.models.client import Client

    c = Client(nev="Árajánlat Teszt (demó) Kft.")
    db.add(c)
    db.flush()
    return c


def _sablon(db, kulcs):
    from app.models.quote import QuoteTemplate

    return db.scalar(select(QuoteTemplate).where(QuoteTemplate.seed_key == kulcs))


def test_seed_idempotens_98_tetel_10_sablon_megjegyzes(db):
    from sqlalchemy import func

    from app.models.quote import QuoteCatalogItem, QuoteNotePreset, QuoteTemplate
    from app.quotes.seed import seed

    seed(db)
    masodik = seed(db)
    assert masodik == {"kategoria": 0, "tetel": 0, "sablon": 0, "megjegyzes": 0}
    assert db.scalar(select(func.count()).select_from(QuoteCatalogItem).where(QuoteCatalogItem.seed_key.is_not(None))) == 98
    assert db.scalar(select(func.count()).select_from(QuoteTemplate).where(QuoteTemplate.seed_key.is_not(None))) == 10
    assert db.scalar(select(func.count()).select_from(QuoteNotePreset)) >= 1
    led = db.scalar(select(QuoteCatalogItem).where(QuoteCatalogItem.seed_key == "c064"))
    assert led.base_price == 720850 and "variant:ledfal" in led.tags and "\n" in led.default_description


def test_seed_nem_irja_felul_a_feluleten_atirt_arat(db):
    from app.models.quote import QuoteCatalogItem
    from app.quotes.seed import seed

    seed(db)
    op = db.scalar(select(QuoteCatalogItem).where(QuoteCatalogItem.seed_key == "c001"))
    op.base_price = 95000
    db.flush()
    seed(db)
    db.refresh(op)
    assert op.base_price == 95000


def test_szamozas_markankent_evente_ujraindul(db):
    from app.quotes.service import kovetkezo_szam

    a = kovetkezo_szam(db, "HYPE", 2091)
    b = kovetkezo_szam(db, "HYPE", 2091)
    c = kovetkezo_szam(db, "CB", 2091)
    assert a == "HYPE-2091-0001" and b == "HYPE-2091-0002" and c == "CB-2091-0001"


def test_sablonbol_letrehozas_opcionalis_sorokkal_es_megjegyzessel(alap, admin, ugyfel):
    from app.quotes import service as S

    t = _sablon(alap, "T1")
    q = S.letrehoz(alap, admin, {"template_id": t.id, "client_id": ugyfel.id, "project_name": "Gála (demó)"})
    assert q.number.startswith("HYPE-") and q.version == 1 and q.status == "draft"
    assert q.summary_label == "A PROJECT TELJES KÖLTSÉGE" and "írásos megrendelő" in q.note_text
    nevek = [(i.name, i.is_optional) for i in q.items]
    assert ("Operatőr", False) in nevek and ("DJI Mavic drón szett", True) in nevek
    gy = next(i for i in q.items if i.name == "Gyártási költség")
    assert gy.unit_price == 30000 and gy.section == "Egyéb"
    # operatőr 80k + lumix 30k + gimbal 7,5k + mikroport 10k + aftermovie 90k + gyártási 30k
    assert q.net_total == 247500
    fotos = next(i for i in q.items if i.name.startswith("Fotós"))
    S.sor_modosit(alap, q, fotos.id, {"is_optional": False})
    assert q.net_total == 347500


def test_katalogusbol_ar_snapshot_es_frissites(alap, admin):
    from app.models.quote import QuoteCatalogItem, QuoteCatalogPriceHistory
    from app.quotes import service as S

    q = S.letrehoz(alap, admin, {"brand": "CB", "project_name": "Üres (demó)"})
    assert q.number.startswith("CB-") and q.items == []
    c = alap.scalar(select(QuoteCatalogItem).where(QuoteCatalogItem.seed_key == "c010"))
    it = S.katalogusbol(alap, q, c.id, "Emberi erőforrás")
    assert it.unit_price == c.base_price == 85000 and it.occasions == 1 and it.quantity == 1
    S.katalogus_ar(alap, c, 95000, admin)
    alap.flush()
    alap.refresh(it)
    assert it.unit_price == 85000  # a meglévő ajánlat nem változik
    tort = alap.scalars(select(QuoteCatalogPriceHistory).where(QuoteCatalogPriceHistory.catalog_item_id == c.id)).all()
    assert [(h.old_price, h.new_price) for h in tort][-1] == (85000, 95000)
    assert S.arak_frissitese(alap, q, [it.id]) == 1 and it.unit_price == 95000 and q.net_total == 95000


def test_beszuras_szekcio_vegere_atrendezes_duplikalas_torles(alap, admin):
    from app.quotes import service as S

    q = S.letrehoz(alap, admin, {"template_id": _sablon(alap, "T4").id, "project_name": "Előadás (demó)"})
    eszkoz = [i for i in q.items if i.section == "Eszközök"]
    uj = S.sor_hozzaad(alap, q, {"section": "Eszközök", "name": "Egyedi sor", "unit_price": 1000})
    rendezett = sorted(q.items, key=lambda x: x.sort_order)
    assert rendezett.index(uj) == rendezett.index(eszkoz[-1]) + 1
    masolat = S.sor_duplikal(alap, q, uj.id)
    assert masolat.name == "Egyedi sor" and masolat.sort_order == uj.sort_order + 10
    ids = [i.id for i in sorted(q.items, key=lambda x: x.sort_order)][::-1]
    S.atrendez(alap, q, [{"id": ids[0], "section": "Első"}] + [{"id": i} for i in ids[1:]])
    elso = min(q.items, key=lambda x: x.sort_order)
    assert elso.id == ids[0] and elso.section == "Első"
    with pytest.raises(S.AjanlatHiba):
        S.atrendez(alap, q, [{"id": ids[0]}])
    S.sor_torol(alap, q, masolat.id)
    assert masolat not in q.items


def test_variansos_csere_katalogustetelre(alap, admin):
    from app.models.quote import QuoteCatalogItem
    from app.quotes import service as S

    q = S.letrehoz(alap, admin, {"template_id": _sablon(alap, "T5").id, "project_name": "LED (demó)"})
    led = next(i for i in q.items if "LED fal" in i.name)
    p39 = alap.scalar(select(QuoteCatalogItem).where(QuoteCatalogItem.seed_key == "c064"))
    S.sor_modosit(alap, q, led.id, {"catalog_item_id": p39.id})
    assert led.name == "15 m² LED fal (P3.9)" and led.unit_price == 720850
    assert q.occasions_label == "Nap"


def test_variant_uj_verzio_es_sablon_mentes(alap, admin, ugyfel):
    from app.quotes import service as S

    q = S.letrehoz(alap, admin, {"template_id": _sablon(alap, "T3").id, "client_id": ugyfel.id,
                                 "project_name": "Stream (demó)", "discount_percent": 5})
    v = S.duplikal(alap, admin, q)
    assert v.number != q.number and v.parent_quote_id == q.id and len(v.items) == len(q.items)
    assert v.net_total == q.net_total and v.client_id == ugyfel.id
    v2 = S.uj_verzio(alap, admin, q)
    v3 = S.uj_verzio(alap, admin, q)
    assert v2.number == q.number and (v2.version, v3.version) == (2, 3)
    op = next(i for i in q.items if i.name == "Operatőr")
    S.sor_modosit(alap, q, op.id, {"unit_price": 90000, "name": "Operatőr (saját)"})
    t = S.sablon_ajanlatbol(alap, q, "Saját stream (demó)")
    ts = next(x for x in t.items if x.catalog_item_id == op.catalog_item_id)
    assert ts.price_override == 90000 and ts.name_override == "Operatőr (saját)"
    tobbi = [x for x in t.items if x.catalog_item_id and x is not ts]
    assert all(x.price_override is None or x.catalog_item.name == "Gyártási költség" for x in tobbi)
    q2 = S.letrehoz(alap, admin, {"template_id": t.id, "project_name": "Újra (demó)"})
    assert sorted((i.name, i.unit_price, i.is_optional) for i in q2.items) == sorted(
        (i.name, i.unit_price, i.is_optional) for i in q.items)


def test_havidijas_sablon_es_utolso_ar_ugyfelnek(alap, admin, ugyfel):
    from app.quotes import service as S

    q = S.letrehoz(alap, admin, {"template_id": _sablon(alap, "T9").id, "client_id": ugyfel.id,
                                 "project_name": "Éves (demó)"})
    o = S.osszegzes(q)
    assert q.pricing_mode == "monthly" and q.months == 12 and o.havidij == round(o.netto / 12)
    pm = next(i for i in q.items if i.name == "Project management havidíj")
    S.sor_modosit(alap, q, pm.id, {"unit_price": 280000})
    tipp = S.utolso_ar(alap, [pm.catalog_item_id], ugyfel.id)
    assert tipp[pm.catalog_item_id]["unit_price"] == 280000


def test_ervenytelen_bemenetek(alap, admin):
    from app.quotes import service as S

    with pytest.raises(S.AjanlatHiba):
        S.letrehoz(alap, admin, {"brand": "XYZ"})
    q = S.letrehoz(alap, admin, {"project_name": "Hibás (demó)"})
    with pytest.raises(S.AjanlatHiba):
        S.modosit(alap, q, {"discount_percent": 120})
    with pytest.raises(S.AjanlatHiba):
        S.sor_hozzaad(alap, q, {"name": "x", "unit": "liter"})
    with pytest.raises(S.AjanlatHiba):
        S.sor_hozzaad(alap, q, {"name": "x", "unit_price": -5})
    S.modosit(alap, q, {"status": "sent"})
    assert q.sent_at is not None


# ── Export ───────────────────────────────────────────────────────────────────


def test_xlsx_export_kepletekkel_opcios_blokkal_es_fajlnev(alap, admin, ugyfel):
    from openpyxl import load_workbook

    from app.quotes import export as E
    from app.quotes import service as S

    q = S.letrehoz(alap, admin, {"template_id": _sablon(alap, "T1").id, "client_id": ugyfel.id,
                                 "project_name": "Gála est (demó)"})
    tartalom = E.xlsx(alap, q)
    ws = load_workbook(io.BytesIO(tartalom)).active
    cellak = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    ertekek = list(cellak.values())
    assert "PROJECT NEVE: Gála est (demó)" in ertekek
    fej = next(r for (r, c), v in cellak.items() if v == "Tétel/Szolgáltatás")
    assert [cellak.get((fej, c)) for c in range(1, 6)] == ["Tétel/Szolgáltatás", "Alkalom", "Mennyiség", "Egységár", "Teljes ár"]
    kepletek = [v for v in ertekek if isinstance(v, str) and v.startswith("=")]
    assert any(v.startswith("=B") and "*C" in v and "*D" in v for v in kepletek)
    assert any(v.startswith("=SUM(") for v in kepletek)
    assert any(isinstance(v, str) and v.startswith("A PROJECT TELJES KÖLTSÉGE") for v in ertekek)
    assert "Opcionális tételek" in ertekek
    assert any(isinstance(v, str) and v.startswith("Megjegyzés:") for v in ertekek)
    assert E.fajlnev(alap, q, "xlsx") == "HYPE_ÁRAJÁNLAT_ÁRAJÁNLAT TESZT (DEMÓ) KFT. - GÁLA EST (DEMÓ).xlsx"


def _kiertekel(ws, cella: str):
    """Az export képleteinek (szorzat, ROUND, SUM, MAX, összeg) kiértékelése -
    így ellenőrizhető, hogy az Excel ugyanazt számolja, mint a szerver."""
    import re

    ertek = ws[cella].value
    if not (isinstance(ertek, str) and ertek.startswith("=")):
        return ertek or 0
    kifejezes = re.sub(r"\b([A-E])(\d+)\b", lambda m: f"v('{m.group(1)}{m.group(2)}')", ertek[1:])
    kornyezet = {
        "v": lambda c: _kiertekel(ws, c), "SUM": lambda *a: sum(a), "MAX": max,
        "ROUND": lambda x, n=0: int(Decimal(str(x)).quantize(Decimal(1), rounding=ROUND_HALF_UP)),
    }
    return eval(kifejezes, {"__builtins__": {}}, kornyezet)  # noqa: S307 - saját, tesztbeli képlet


def test_xlsx_kepletek_ugyanazt_szamoljak_mint_a_szerver(alap, admin):
    from openpyxl import load_workbook

    from app.quotes import export as E
    from app.quotes import service as S

    q = S.letrehoz(alap, admin, {"template_id": _sablon(alap, "T3").id, "project_name": "Képlet (demó)"})
    S.modosit(alap, q, {"discount_percent": 7.5, "discount_amount": 1000})
    S.sor_modosit(alap, q, q.items[0].id, {"line_discount_percent": 12.5, "occasions": 1.5})
    ws = load_workbook(io.BytesIO(E.xlsx(alap, q))).active
    sor = next(c.row for c in ws["A"] if c.value == "STREAMING SZOLGÁLTATÁS KÖLTSÉGE")
    assert _kiertekel(ws, f"E{sor}") == S.osszegzes(q).netto == q.net_total


def test_xlsx_havidijas_osszesito_sorok(alap, admin):
    from openpyxl import load_workbook

    from app.quotes import export as E
    from app.quotes import service as S

    q = S.letrehoz(alap, admin, {"template_id": _sablon(alap, "T9").id, "project_name": "Keret (demó)"})
    ws = load_workbook(io.BytesIO(E.xlsx(alap, q))).active
    ertekek = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "A PROJECT HAVIDÍJA" in ertekek and "A PROJECT TELJES KÖLTSÉGE 12 HÓNAPRA" in ertekek
    assert "Teljes ár (12 hónap)" in ertekek


def test_pdf_export_ervenyes_pdf(alap, admin):
    from pypdf import PdfReader

    from app.quotes import export as E
    from app.quotes import service as S

    q = S.letrehoz(alap, admin, {"template_id": _sablon(alap, "T3").id, "project_name": "Közvetítés (demó)"})
    tartalom = E.pdf(alap, q)
    assert tartalom[:5] == b"%PDF-"
    szoveg = "".join(p.extract_text() for p in PdfReader(io.BytesIO(tartalom)).pages)
    assert "Közvetítés (demó)" in szoveg and "STREAMING SZOLGÁLTATÁS KÖLTSÉGE" in szoveg
    assert "Hype Productions Kft." in szoveg


# ── API ──────────────────────────────────────────────────────────────────────


@pytest.fixture()
def kliens(db, admin, monkeypatch):
    from fastapi.testclient import TestClient

    from app.core.database import get_db
    from app.core.security import get_current_user
    from app.main import app

    monkeypatch.setattr(db, "commit", db.flush)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: admin
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()


def test_api_sablonbol_letrehozas_szerkesztes_export(kliens, alap, ugyfel):
    t = _sablon(alap, "T1")
    sablonok = kliens.get("/api/v1/quotes/templates").json()
    kartya = next(x for x in sablonok if x["id"] == t.id)
    assert kartya["typical_total"] == 700000 and kartya["base_total"] == 247500 and kartya["item_count"] >= 10

    r = kliens.post("/api/v1/quotes", json={"template_id": t.id, "brand": "HYPE", "client_id": ugyfel.id,
                                            "project_name": "API gála (demó)"})
    assert r.status_code == 201, r.text
    q = r.json()
    assert q["totals"]["netto"] == 247500 and q["client"]["nev"] == ugyfel.nev
    op = next(i for i in q["items"] if i["name"] == "Operatőr")
    assert op["catalog_price"] == 80000

    r = kliens.patch(f"/api/v1/quotes/{q['id']}/items/{op['id']}", json={"quantity": 2})
    assert r.status_code == 200 and r.json()["totals"]["netto"] == 327500

    kat = kliens.get("/api/v1/quotes/catalog/items", params={"q": "gimbal"}).json()
    assert kat and kat[0]["name"] == "Kameramozgató (gimbal)"
    r = kliens.post(f"/api/v1/quotes/{q['id']}/items/from-catalog",
                    json={"catalog_item_id": kat[0]["id"], "section": "Eszközök"})
    assert r.status_code == 201 and r.json()["totals"]["netto"] == 335000

    lista = kliens.get("/api/v1/quotes", params={"q": "API gála"}).json()
    assert [x["id"] for x in lista] == [q["id"]] and lista[0]["net_total"] == 335000

    x = kliens.get(f"/api/v1/quotes/{q['id']}/export.xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert "attachment" in x.headers["content-disposition"] and "filename*=UTF-8''HYPE_" in x.headers["content-disposition"]
    p = kliens.get(f"/api/v1/quotes/{q['id']}/export.pdf")
    assert p.status_code == 200 and p.content[:5] == b"%PDF-"

    r = kliens.post(f"/api/v1/quotes/{q['id']}/new-version")
    assert r.status_code == 201 and r.json()["version"] == 2 and r.json()["number"] == q["number"]
    verziok = kliens.get(f"/api/v1/quotes/{q['id']}").json()["versions"]
    assert [v["version"] for v in verziok] == [1, 2]


def test_api_katalogus_arvaltozas_tortenet_es_jogosultsag_nelkuli_404(kliens, alap):
    c = kliens.get("/api/v1/quotes/catalog/items", params={"q": "Hangmérnök"}).json()[0]
    r = kliens.patch(f"/api/v1/quotes/catalog/items/{c['id']}", json={"base_price": 80000})
    assert r.status_code == 200 and r.json()["base_price"] == 80000
    tort = kliens.get(f"/api/v1/quotes/catalog/items/{c['id']}/history").json()
    assert tort[0]["old_price"] == 75000 and tort[0]["new_price"] == 80000
    assert kliens.get("/api/v1/quotes/999999999").status_code == 404
    assert kliens.patch("/api/v1/quotes/catalog/items/999999999", json={"base_price": 1}).status_code == 404
